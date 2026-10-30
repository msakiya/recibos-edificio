"""Datos REALES de setiembre (fuera del repo) + Drive simulado:
sin ZIP, PDF por dpto, Excel actualizado, subida a Drive (carpetas, reemplazo, link público),
correo (HTML con VER RECIBO, cuenta, titular, Importante), WhatsApp (URL y 'no disponible'),
Desconectar y conectar otra cuenta. Capturas con datos reales → gitignored (captura-*.png)."""
import time, subprocess, signal, urllib.parse, sys
from pathlib import Path
import openpyxl
from playwright.sync_api import sync_playwright
sys.path.insert(0, str(Path(__file__).parent))
from mock_google import DriveFalso, instalar

ROOT = Path('/workspace/recibos_web')
TEST = Path('/workspace/recibos_web_test')
FOTOS = TEST / 'fotos'
SET = Path('/workspace/sistema_recibos/set_2026')
URL = 'http://127.0.0.1:8765/'

def config_real():
    """Datos de pago leídos del Excel real de prueba (para no escribir datos reales en el repo público)."""
    wb = openpyxl.load_workbook(TEST / 'setiembre-real-envio.xlsx')
    cfg = {r[0].value: r[1].value for r in wb['Config'].iter_rows(min_row=2) if r[0].value}
    tel401 = [r[2].value for r in wb['Departamentos'].iter_rows(min_row=2) if r[0].value == 401][0]
    return cfg, ''.join(ch for ch in str(tel401) if ch.isdigit())

def generar(page):
    page.click('#btnGenerar')
    try:
        page.wait_for_selector('#btnSiGen', timeout=5000); page.click('#btnSiGen')
    except Exception:
        pass

def main():
    srv = subprocess.Popen(['python3', '-m', 'http.server', '8765', '--bind', '127.0.0.1'],
                           cwd=str(ROOT), stdout=subprocess.DEVNULL, stderr=subprocess.DEVNULL)
    time.sleep(0.6)
    drive = DriveFalso()
    cfg_r, tel401 = config_real()
    CUENTA, CCI, TITULAR, WA = cfg_r['CUENTA_BANCARIA'], cfg_r['CCI'], cfg_r['TITULAR_CUENTA'], cfg_r['WHATSAPP_CONSTANCIA']
    try:
        with sync_playwright() as p:
            browser = p.chromium.launch(executable_path='/usr/bin/google-chrome', args=['--no-sandbox'])
            ctx = browser.new_context(viewport={'width': 1100, 'height': 900}, accept_downloads=True)
            ctx.grant_permissions(['clipboard-read', 'clipboard-write'], origin=URL.rstrip('/'))
            instalar(ctx, drive)
            page = ctx.new_page()
            errs, descargas = [], []
            page.on('pageerror', lambda e: errs.append(str(e)))
            page.on('download', lambda d: descargas.append(d.suggested_filename))
            page.on('dialog', lambda d: d.accept())
            page.goto(URL, wait_until='networkidle')
            page.evaluate("() => { window.__abiertas = []; window.open = (u) => { window.__abiertas.push(u); return {}; }; }")

            # Conectar Drive (un botón)
            page.wait_for_selector('#btnConectar:not([disabled])')
            page.click('#btnConectar')
            page.wait_for_selector('text=Google Drive conectado')
            assert 'tesoreria.prueba@gmail.com' in page.inner_text('#barraDrive0')
            req = page.evaluate('() => window.__gis.reqs[0]'); cfg = page.evaluate('() => window.__gis.cfg')
            assert req.get('prompt') == 'select_account', req
            assert cfg['scope'] == 'https://www.googleapis.com/auth/drive.file', cfg
            ls = page.evaluate("() => JSON.parse(localStorage.getItem('recibos.drive'))")
            assert ls == {'conectado': True, 'email': 'tesoreria.prueba@gmail.com', 'nombre': 'Prueba'}, ls
            assert 'tok-' not in str(page.evaluate("() => JSON.stringify(localStorage)"))  # el token no se guarda

            # Recarga: sigue "conectado" (localStorage) y pide token silencioso al generar
            page.reload(wait_until='networkidle')
            page.evaluate("() => { window.__abiertas = []; window.open = (u) => { window.__abiertas.push(u); return {}; }; }")
            assert 'tesoreria.prueba@gmail.com' in page.inner_text('#barraDrive0')

            page.set_input_files('#fileExcel', str(TEST / 'setiembre-real-envio.xlsx'))
            page.wait_for_selector('#btnAPaso3b', timeout=15000)
            page.click('#btnAPaso3b')
            fotos = sorted([str(f) for f in FOTOS.iterdir() if f.suffix.lower() in ('.jpg', '.jpeg', '.png')])
            fotos += [str(SET / 'recibo-de-agua-2026.09.19.jpg'), str(SET / 'Luz_del_sur_Setiembre.pdf')]
            page.set_input_files('#fileFotos', fotos)
            page.wait_for_function("() => window.__RecibosApp.estado.cola.length && window.__RecibosApp.estado.cola.every(c => c.status !== 'procesando')", timeout=120000)
            page.wait_for_function("() => !!window.__RecibosApp.estado.archivos.agua && !!window.__RecibosApp.estado.archivos.luz", timeout=30000)
            assert 'ZIP' not in page.inner_text('#btnGenerar')
            generar(page)
            page.wait_for_selector('.envio-fila', timeout=300000)
            page.wait_for_selector('text=Los 10 recibos están en Google Drive', timeout=60000)
            assert descargas == [], descargas            # sin ZIP ni descargas automáticas
            req2 = page.evaluate('() => window.__gis.reqs[0]')  # página recargada
            assert req2.get('prompt') == '' and req2.get('hint') == 'tesoreria.prueba@gmail.com', req2

            # Drive: carpetas y archivos
            rutas = sorted(drive.ruta(f) for f in drive.pdfs())
            print('DRIVE:', rutas[0], '...', len(rutas))
            assert len(rutas) == 10
            assert rutas[0] == 'Recibos Edificio Nápoles II/2026/9-Setiembre/RECIBOS/Recibo - 101 - Setiembre 2026.pdf', rutas[0]
            assert len(drive.carpetas()) == 4
            assert all(drive.perms[f['id']] == [{'role': 'reader', 'type': 'anyone'}] for f in drive.pdfs())

            filas = page.locator('.envio-fila')
            assert filas.count() == 10
            f101 = filas.nth(0)
            assert 'Dpto 101' in f101.inner_text() and 'S/ 471.99' in f101.inner_text(), f101.inner_text()
            id101 = [f['id'] for f in drive.pdfs() if f['name'].startswith('Recibo - 101 ')][0]
            link101 = 'https://drive.google.com/file/d/%s/view?usp=drivesdk' % id101

            # PDF individual
            with page.expect_download() as d:
                f101.locator('.envio-pdf').click()
            assert d.value.suggested_filename == 'Recibo - 101 - Setiembre 2026.pdf', d.value.suggested_filename
            assert d.value.path().read_bytes()[:4] == b'%PDF'
            # Excel actualizado
            with page.expect_download() as dx:
                page.click('#btnExcel')
            assert dx.value.suggested_filename == 'Excel actualizado - para Octubre 2026.xlsx', dx.value.suggested_filename
            px = TEST / 'excel-actualizado-envio.xlsx'; dx.value.save_as(str(px))
            wbx = openpyxl.load_workbook(px)
            assert wbx['Mes a facturar']['B4'].value == '2026-10'

            # Correo 101
            f101.locator('.envio-btn').click()
            page.wait_for_selector('#toast.visible:has-text("Ctrl+V")', timeout=5000)
            toast = page.inner_text('#toast'); print('TOAST:', toast)
            assert 'Ctrl+V' in toast and 'Mensaje copiado' in toast
            assert f101.locator('.envio-btn').inner_text().strip().endswith('Enviado ✓')
            url = page.evaluate('() => window.__abiertas[0]')
            q = urllib.parse.parse_qs(urllib.parse.urlparse(url).query)
            assert url.startswith('https://mail.google.com/mail/?view=cm&fs=1&')
            assert q['to'] == ['dpto101@prueba.test'] and q['su'] == ['Recibo de Gastos Comunes - Setiembre 2026 - Dpto 101'], q
            clip = page.evaluate("""async () => { const items = await navigator.clipboard.read(); const out = {};
                for (const it of items) for (const t of it.types) out[t] = await (await it.getType(t)).text(); return out; }""")
            html, txt = clip.get('text/html', ''), clip.get('text/plain', '')
            for s in ['Edificio Nápoles II', 'Hola, (Dpto 101)', 'correspondiente al mes de <b>Setiembre 2026</b>.',
                      'Total a pagar: S/ 471.99', '>VER RECIBO</a>', 'href="%s"' % link101, '<b>Cuenta BCP:</b> ' + CUENTA,
                      '<b>CCI:</b> ' + CCI, '<b>Titular de la Cuenta:</b> ' + TITULAR,
                      '<b>Importante:</b> Enviar constancia de pago al WhatsApp: <b>' + WA + '</b>']:
                assert s in html, s
            assert 'VER RECIBO: ' + link101 in txt
            # Orden como la imagen
            orden = [html.index(s) for s in ['Nápoles II', 'Hola,', 'border-top:2px', 'Se ha generado', 'Total a pagar', 'VER RECIBO', 'Cuenta BCP', 'CCI:', 'Titular', '#cccccc', 'Importante']]
            assert orden == sorted(orden), orden

            # WhatsApp: 101 sin teléfono → no disponible; 401 con teléfono
            wa101 = f101.locator('.wa-btn')
            assert wa101.is_disabled() and wa101.inner_text().strip() == 'WhatsApp no disponible'
            assert wa101.get_attribute('title') == 'Sin teléfono en el Excel'
            f401 = filas.nth(6); assert 'Dpto 401' in f401.inner_text()
            f401.locator('.wa-btn').click()
            wurl = page.evaluate('() => window.__abiertas[window.__abiertas.length-1]')
            
            assert wurl.startswith('https://web.whatsapp.com/send?phone=51' + tel401 + '&text='), wurl
            msg = urllib.parse.parse_qs(urllib.parse.urlparse(wurl).query)['text'][0]
            
            id401 = [f['id'] for f in drive.pdfs() if f['name'].startswith('Recibo - 401 ')][0]
            for s in ['Dpto 401', '*Setiembre 2026*', '*Total a pagar: S/ ', 'Puedes ver tu recibo aquí 👉 https://drive.google.com/file/d/%s/view' % id401,
                      '*Cuenta BCP:* ' + CUENTA, '*CCI:* ' + CCI, '*Titular de la Cuenta:* ' + TITULAR,
                      'Importante:* Enviar constancia de pago al WhatsApp: *' + WA + '*']:
                assert s in msg, s
            assert f401.locator('.wa-btn').inner_text().strip().endswith('Enviado ✓')
            assert 'Puedes ver tu recibo aquí' in page.inner_text('#prevWa')

            # Capturas (datos reales → no van al repo)
            page.evaluate("() => document.getElementById('toast').classList.remove('visible')")
            page.wait_for_timeout(400)
            page.locator('#panelResultado').screenshot(path=str(ROOT / 'captura-resultado-drive.png'))
            page.set_viewport_size({'width': 390, 'height': 844}); page.wait_for_timeout(300)
            page.locator('#panelResultado').screenshot(path=str(ROOT / 'captura-resultado-movil.png'))
            page.set_viewport_size({'width': 1100, 'height': 900})

            # Generar otra vez: reutiliza carpetas y reemplaza los PDF (mismo nombre)
            generar(page)
            page.wait_for_selector('text=Los 10 recibos están en Google Drive', timeout=300000)
            page.wait_for_timeout(300)
            assert len(drive.carpetas()) == 4 and len(drive.pdfs()) == 10, (len(drive.carpetas()), len(drive.pdfs()))
            assert all(f.get('versiones') == 2 for f in drive.pdfs())
            assert descargas == ['Recibo - 101 - Setiembre 2026.pdf', 'Excel actualizado - para Octubre 2026.xlsx'], descargas

            # Desconectar → revoca y borra; luego conectar otra cuenta
            tok = page.evaluate('() => "tok-" + window.__gis.n')
            page.click('#btnDesconectar')
            page.wait_for_selector('#btnConectar')
            assert page.evaluate('() => window.__gis.revoked') == [tok], page.evaluate('() => window.__gis.revoked')
            assert page.evaluate("() => localStorage.getItem('recibos.drive')") is None
            drive.emails[0] = 'otra.cuenta@gmail.com'
            page.click('#btnConectar')
            page.wait_for_selector('text=otra.cuenta@gmail.com')
            assert page.evaluate("() => JSON.parse(localStorage.getItem('recibos.drive')).email") == 'otra.cuenta@gmail.com'

            assert not errs, errs
            browser.close()
            print('ENVIO + DRIVE OK')
    finally:
        srv.send_signal(signal.SIGTERM); srv.wait(timeout=5)

if __name__ == '__main__':
    main()
