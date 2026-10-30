"""Prueba de punta a punta: carga Excel real + fotos + genera los 10 PDF (sin ZIP) y el Excel actualizado."""
import os, re, time, subprocess, signal
import openpyxl
from pathlib import Path
from playwright.sync_api import sync_playwright

ROOT = Path('/workspace/recibos_web')
OUT = ROOT  # capturas y previews aquí
TEST = Path('/workspace/recibos_web_test')
FOTOS = TEST / 'fotos'
SET = Path('/workspace/sistema_recibos/set_2026')

def main():
    # Servidor estático
    srv = subprocess.Popen(['python3', '-m', 'http.server', '8765', '--bind', '127.0.0.1'],
                           cwd=str(ROOT), stdout=subprocess.DEVNULL, stderr=subprocess.DEVNULL)
    time.sleep(0.6)
    try:
        with sync_playwright() as p:
            browser = p.chromium.launch(executable_path='/usr/bin/google-chrome',
                                        args=['--no-sandbox', '--disable-web-security'])
            ctx = browser.new_context(viewport={'width': 1100, 'height': 900}, accept_downloads=True)
            page = ctx.new_page()
            errs = []
            page.on('pageerror', lambda e: errs.append(str(e)))
            page.goto('http://127.0.0.1:8765/', wait_until='networkidle')
            page.screenshot(path=str(OUT / 'captura-paso1.png'), full_page=True)

            # Excel antiguo (sin pestaña "Mes a facturar" y con setiembre ya generado): pide la pestaña, no campos a mano
            page.set_input_files('#fileExcel', str(TEST / 'setiembre-real.xlsx'))
            page.wait_for_selector('text=Falta completar el Excel', timeout=15000)
            assert page.locator('#btnConHoja').count() == 1
            with page.expect_download() as dh:
                page.click('#btnConHoja')
            p_conhoja = TEST / 'antiguo-con-hoja.xlsx'; dh.value.save_as(str(p_conhoja))
            wbh = openpyxl.load_workbook(p_conhoja)
            assert wbh.sheetnames[0] == 'Mes a facturar' and wbh['Mes a facturar']['B4'].value == '2026-10', wbh.sheetnames
            page.screenshot(path=str(OUT / 'captura-excel-antiguo.png'), full_page=True)

            # Excel real con la pestaña "Mes a facturar" = Setiembre: resumen de solo lectura, sin paso 2
            page.set_input_files('#fileExcel', str(TEST / 'setiembre-real-mes.xlsx'))
            page.wait_for_selector('#btnAPaso3b', timeout=15000)
            assert page.locator('#paso2').count() == 0 and page.locator('#aguaImporte').count() == 0
            txt = page.inner_text('#excelInfo')
            for t in ['Setiembre 2026', 'S/ 679.40', '154 m³', 'S/ 510.40', 'S/ 471.99', 'S/ 597.80', 'S/ 538.00',
                      'S/ 4,987.48', 'S/ 4,897.67', '1,281.328', '1,290.107', '25/09/2026', '05/10/2026']:
                assert t in txt, t
            page.screenshot(path=str(OUT / 'captura-resumen.png'), full_page=True)
            page.click('#btnAPaso3b')
            page.wait_for_selector('#paso3:not(.oculto)')
            page.wait_for_function("() => document.getElementById('paso1').classList.contains('oculto')")

            # Fotos: medidores + comprobantes + recibos
            fotos = sorted([str(f) for f in FOTOS.iterdir() if f.suffix.lower() in ('.jpg', '.jpeg', '.png')])
            fotos += [str(SET / 'recibo-de-agua-2026.09.19.jpg'), str(SET / 'Luz_del_sur_Setiembre.pdf')]
            page.set_input_files('#fileFotos', fotos)
            page.wait_for_function(
                "() => window.__RecibosApp.estado.cola.length && window.__RecibosApp.estado.cola.every(c => c.status !== 'procesando')",
                timeout=120000)
            # Espera a que agua y luz estén asignados
            page.wait_for_function(
                "() => !!window.__RecibosApp.estado.archivos.agua && !!window.__RecibosApp.estado.archivos.luz",
                timeout=30000)
            page.screenshot(path=str(OUT / 'captura-paso3.png'), full_page=True)
            n_med = page.evaluate("() => Object.keys(window.__RecibosApp.estado.archivos.medidor).filter(k => !k.includes('__')).length")
            print('medidores asignados:', n_med)
            assert n_med >= 9, n_med

            # Generar (sin ZIP: lista por dpto con ⬇ PDF y botón de Excel actualizado)
            page.click('#btnGenerar')
            try:
                page.wait_for_selector('#btnSiGen', timeout=5000)
                page.click('#btnSiGen')
            except Exception:
                pass
            page.wait_for_selector('text=Listo:', timeout=300000)
            page.screenshot(path=str(OUT / 'captura-resultado.png'), full_page=True)
            assert page.locator('.envio-pdf').count() == 10
            with page.expect_download() as dp:
                page.locator('.envio-pdf').nth(0).click()
            assert dp.value.suggested_filename == 'Recibo - 101 - Setiembre 2026.pdf'
            dp.value.save_as(str(TEST / 'recibo-101.pdf'))
            assert (TEST / 'recibo-101.pdf').read_bytes()[:4] == b'%PDF'
            with page.expect_download() as dx:
                page.click('#btnExcel')
            dx.value.save_as(str(TEST / 'actualizado-setiembre.xlsx'))

            # Excel actualizado: pestaña lista para Octubre
            wbo = openpyxl.load_workbook(TEST / 'actualizado-setiembre.xlsx')
            ws = wbo['Mes a facturar']
            assert wbo.sheetnames[0] == 'Mes a facturar'
            assert ws['B4'].value == '2026-10', ws['B4'].value
            assert ws['B5'].value is None and ws['C16'].value == 1290.107 and ws['D16'].value is None
            assert abs(ws['B10'].value - 4897.67) < 1e-9
            # Ella llena Octubre en Excel (valores de su archivo "Octubre - 2026.xlsx"), sube y revisa
            ws['B5'] = 632.6; ws['B6'] = 144; ws['B7'] = 532.3
            act = [1290.107, 1752.792, 53.561, 697.997, 469.849, 1008.439, 1162.571, 297.992, 1843.492, 786.497]
            for i, v in enumerate(act): ws.cell(row=16 + i, column=4, value=v)
            p_oct = TEST / 'octubre-lleno.xlsx'; wbo.save(p_oct)
            page.goto('http://127.0.0.1:8765/', wait_until='networkidle')
            page.set_input_files('#fileExcel', str(p_oct))
            page.wait_for_selector('#btnAPaso3b', timeout=15000)
            txt = page.inner_text('#excelInfo')
            for t in ['Octubre 2026', 'S/ 632.60', '144 m³', 'S/ 532.30', '1,752.792']:
                assert t in txt, t
            page.screenshot(path=str(OUT / 'captura-resumen-octubre.png'), full_page=True)
            # Falta una lectura → error con la celda exacta
            ws['D18'].value = None; wbo.save(TEST / 'octubre-falta.xlsx')
            page.set_input_files('#fileExcel', str(TEST / 'octubre-falta.xlsx'))
            page.wait_for_selector('text=Falta completar el Excel', timeout=15000)
            err = page.inner_text('#excelInfo'); print('ERROR MOSTRADO:', err[:400])
            assert '201 (D18)' in err and 'Mes a facturar' in err
            assert page.locator('#btnAPaso3b').count() == 0
            page.screenshot(path=str(OUT / 'captura-error-celda.png'), full_page=True)

            assert not errs, errs
            browser.close()
    finally:
        srv.send_signal(signal.SIGTERM)
        srv.wait(timeout=5)

    # Renderiza PDF 101 a PNG con pdftoppm
    subprocess.check_call(['pdftoppm', '-png', '-r', '110', str(TEST / 'recibo-101.pdf'), str(OUT / 'preview-101')])
    # Renombra preview-101-1.png → preview-101-p1.png
    for i, src in enumerate(sorted(OUT.glob('preview-101-*.png')), 1):
        if re.search(r'preview-101-\d+\.png$', src.name):
            dest = OUT / f'preview-101-p{i}.png'
            src.rename(dest)
            print('→', dest)
    print('E2E OK')

if __name__ == '__main__':
    main()
